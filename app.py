import logging
import re
import os
import io
import lxml.etree
from flask import Flask, request, render_template, flash, redirect
from werkzeug.utils import secure_filename
from docx import Document
import mammoth
from tabulate import tabulate

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = './uploads'
app.config['ALLOWED_EXTENSIONS'] = {'docx'}

if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def word_table_to_markdown(filename):
    doc = Document(filename)
    markdown_content = []

    for i, table in enumerate(doc.tables):
        data = []

        for row in table.rows:
            cell_data = []

            for cell in row.cells:
                cell_data.append(cell.text)

            data.append(cell_data)

        placeholder = f'[TABLE_{i}]'
        custom_element = lxml.etree.Element('placeholder')
        custom_element.text = placeholder
        table._element.getparent().replace(table._element, custom_element)

        markdown_content.append((placeholder, data))

    return markdown_content

def replace_images_with_placeholder(content, images):
    img_regex = r'!\[.*?\]\((.*?)\)'
    placeholder = '[ADD IMAGE HERE]'
    result = re.sub(img_regex, lambda match: placeholder if match.group(1) in images else match.group(0), content)
    result = re.sub(r'\(data:image.*?\)', '', result)
    return result

def replace_urls(content):
    result = re.sub(r'https://learn.microsoft.com/azure/(.+)', r'/azure/\1', content)
    result = result.replace("/en-us/", "/")
    return result

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('No file part')
            return redirect(request.url)
        file = request.files['file']
        if file.filename == '':
            flash('No selected file')
            return redirect(request.url)
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename).replace("\\", "/")
            file.save(filepath)

            with open(filepath, "rb") as docx_file:
                result = mammoth.convert_to_markdown(docx_file).value

            table_data = word_table_to_markdown(filepath)
            table_placeholders = [data[0] for data in table_data]

            for placeholder, table in table_data:
                if placeholder in result:
                    table_markdown = tabulate(table, tablefmt="pipe", headers="firstrow")
                    result = result.replace(placeholder, table_markdown, 1)

            duplicate_table_markers = re.findall(r"\[TABLE_\d+\]", result)
            for marker in duplicate_table_markers:
                result = result.replace(marker, "")

            markdown_content = [result]
            for _, table in table_data:
                table_markdown = tabulate(table, tablefmt="pipe", headers="firstrow")
                markdown_content.append(table_markdown)

            result = "\n\n".join(markdown_content)

            result = result.rstrip()
            result = result.replace("\\.", ".").replace("\\)", ")").replace("\\-", "-").replace("\\(", "(").replace("\\]", "]").replace("\\[", "[").replace("\\#", "#")
            result = re.sub(r'https://learn.microsoft.com/[a-z]+-[a-z]+(/azure)?', '/azure', result)
            result = result.replace('  ', ' ')
            result = re.sub(r'__(.*?)__', r'**\1**', result)
            result = re.sub(r' \*\*(.*?)', r'** \1', result)
            result = re.sub(r'-\*\*\s(.+)', r'- **\1', result)
            result = result.replace("\\`", "`")
            result = re.sub(r'(\S)\s+\]', r'\1]', result)
            result = re.sub(r'(-{4,})', '---', result)

            images = re.findall(r'!\[.*?\]\((.*?)\)', result)
            result = replace_images_with_placeholder(result, images)

            os.remove(filepath)
            return render_template('index.html', markdown_content=result)

    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True)
